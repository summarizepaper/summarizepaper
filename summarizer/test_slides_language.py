import copy
import io
import json
from unittest import mock

from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation

from . import generated, slides, slides_pptx
from .models import Generated, UserPlan
from .test_tools import ToolsTestCase


def content(lang):
    french = lang == 'fr'
    return {'title': 'Les nuages d’une planète chaude' if french else 'Clouds on a hot planet',
            'slides': [{'title': 'Résultats principaux' if french else 'Main results',
                        'bullets': ['Une atmosphère nuageuse est observée.' if french else 'A cloudy atmosphere is observed.'],
                        'notes': 'Présenter les incertitudes.' if french else 'Discuss the uncertainties.', 'figure': None}],
            'strengths': ['Mesures directes' if french else 'Direct measurements'],
            'limitations': ['Échantillon limité' if french else 'Limited sample'],
            'questions': ['Quelle mesure supplémentaire ?' if french else 'Which further measurement?'],
            'license': 'CC BY 4.0', 'language': lang, 'generation_version': 2}


class SlidesLanguageTests(ToolsTestCase):
    def setUp(self):
        super().setUp()
        self.paper()
        self.reader = self.user('language-reader', paid=True)
        self.client = self.client_for(self.reader)
        self.base = '/en/journal-club/2101.00001v1/'

    def create(self, lang, kind=slides.DECK_KIND):
        return Generated.objects.create(kind=kind, key='2101.00001v1', lang=lang, status='done', content=json.dumps(content(lang)))

    def test_language_can_be_selected_before_any_generation(self):
        response = self.client.get(self.base)
        soup = BeautifulSoup(response.content, 'html.parser')
        form = soup.select_one('#deck-language-form')
        self.assertEqual([option['value'] for option in form.select('option')], ['fr', 'en'])
        self.assertEqual(form.select_one('option[selected]')['value'], 'en')
        french = self.client.get(self.base, {'deck_lang': 'fr'})
        self.assertEqual(french.context['deck_lang'], 'fr')
        self.assertEqual(Generated.objects.count(), 0)

    def test_selected_generation_language_is_independent_from_interface_and_has_its_own_cache(self):
        with mock.patch.object(slides, 'figures', return_value=[]), \
             mock.patch('summarizer.paper_source.paper_text', return_value='Original paper in English.'), \
             mock.patch.object(generated, 'chat_json', return_value=content('fr')) as chat:
            result = self.client.post(self.base + 'start/', {'deck_lang': 'fr'})
            again = self.client.post(self.base + 'start/', {'deck_lang': 'fr'})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['id'], again.json()['id'])
        self.assertEqual(chat.call_count, 1)
        self.assertIn('Every title, bullet, speaker note', chat.call_args.args[0])
        self.assertIn('must be in French', chat.call_args.args[0])
        self.assertIn('deck_lang=fr', result.json()['url'])
        row = Generated.objects.get(kind=slides.DECK_KIND, lang='fr')
        self.assertEqual(row.data()['language'], 'fr')
        self.assertEqual(row.data()['title'], 'Les nuages d’une planète chaude')
        self.assertFalse(Generated.objects.filter(lang='en').exists())
        page = self.client.get(result.json()['url'])
        self.assertEqual(page.context['pages'][0]['title'], 'Les nuages d’une planète chaude')

    def test_slide_labels_pdf_canvas_and_powerpoint_use_the_content_language(self):
        for interface, lang in [('en', 'fr'), ('fr', 'en')]:
            row = self.create(lang)
            data = row.data()
            data['slides'][0]['figure'] = {'number': 1, 'url': 'https://arxiv.org/html/2101.00001v1/a.png', 'caption': 'Original caption'}
            row.content = json.dumps(data)
            row.save()
            url = '/%s/journal-club/2101.00001v1/?deck_lang=%s&deck_version=2' % (interface, lang)
            response = self.client.get(url)
            soup = BeautifulSoup(response.content, 'html.parser')
            canvas = soup.select_one('#deck').get_text(' ', strip=True)
            self.assertIn('Points forts et limites' if lang == 'fr' else 'Strengths and limits', canvas)
            caption = soup.select_one('figcaption').get_text(' ', strip=True)
            self.assertIn("de l'article" if lang == 'fr' else 'of the paper', caption)
            payload = {'revision': response.context['deck_revision'], 'pages': response.context['pages']}
            png = io.BytesIO()
            Image.new('RGB', (200, 100), 'white').save(png, 'PNG')
            with mock.patch.object(slides_pptx, 'figure_bytes', return_value=png.getvalue()):
                exported = self.client.post(response.context['powerpoint_url'], json.dumps(payload), content_type='application/json')
            self.assertEqual(exported.status_code, 200)
            deck = Presentation(io.BytesIO(exported.content))
            self.assertEqual(deck.core_properties.language, lang)
            text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn(data['title'], text)
            self.assertIn('Points forts' if lang == 'fr' else 'Strengths', text)
            self.assertNotIn('Strengths' if lang == 'fr' else 'Points forts', text)
            self.assertIn('de l’article' if lang == 'fr' else 'of the paper', text)
            self.assertEqual(deck.slides[1].notes_slide.notes_text_frame.text, data['slides'][0]['notes'])

    def test_previous_decks_and_private_edits_are_retained_after_new_generation(self):
        legacy = self.create('fr', kind='slides')
        legacy_data = legacy.data()
        legacy_data.pop('generation_version')
        legacy_data.pop('language')
        legacy.content = json.dumps(legacy_data)
        legacy.save()
        response = self.client.get(self.base, {'deck_lang': 'fr', 'deck_version': '1'})
        payload = {'revision': response.context['deck_revision'], 'pages': copy.deepcopy(response.context['pages'])}
        payload['pages'][1]['title'] = 'Ma présentation précédente'
        self.client.post(response.context['save_url'], json.dumps(payload), content_type='application/json')
        old_copy = Generated.objects.get(kind='slides_edit', user=self.reader)
        previous_content = old_copy.content
        self.create('fr')
        current = self.client.get(self.base, {'deck_lang': 'fr'})
        self.assertNotContains(current, 'Ma présentation précédente')
        self.assertTrue(current.context['previous_deck_url'])
        self.assertContains(self.client.get(current.context['previous_deck_url']), 'Ma présentation précédente')
        old_copy.refresh_from_db()
        self.assertEqual(old_copy.content, previous_content)

    def test_language_specific_private_copies_are_separate_and_owner_only(self):
        for lang in ('fr', 'en'):
            self.create(lang)
        page = self.client.get(self.base, {'deck_lang': 'fr'})
        payload = {'revision': page.context['deck_revision'], 'pages': copy.deepcopy(page.context['pages'])}
        payload['pages'][1]['title'] = 'Ma version française'
        self.assertEqual(self.client.post(page.context['save_url'], json.dumps(payload), content_type='application/json').status_code, 200)
        private = Generated.objects.get(kind='slides_edit_v2', lang='fr')
        self.assertNotContains(self.client.get(self.base, {'deck_lang': 'en'}), 'Ma version française')
        other = self.client_for(self.user('another-language-reader', paid=True))
        self.assertNotContains(other.get(self.base, {'deck_lang': 'fr'}), 'Ma version française')
        self.assertEqual(other.get('/en/ai/%s/' % private.pk).status_code, 404)
        UserPlan.objects.filter(user=self.reader).update(paid=False)
        self.assertEqual(self.client.get('/en/ai/%s/' % private.pk).status_code, 404)

    def test_unsupported_languages_and_versions_are_rejected_without_model_calls(self):
        with mock.patch.object(generated, 'chat_json') as chat:
            self.assertEqual(self.client.post(self.base + 'start/', {'deck_lang': 'unsupported'}).status_code, 400)
            self.assertEqual(self.client.get(self.base, {'deck_lang': 'unsupported'}).status_code, 400)
            self.assertEqual(self.client.get(self.base, {'deck_version': '999'}).status_code, 400)
        chat.assert_not_called()
        self.assertFalse(Generated.objects.exists())

    def test_generated_title_is_required_instead_of_reusing_the_original_paper_title(self):
        answer = content('fr')
        answer.pop('title')
        with mock.patch.object(slides, 'figures', return_value=[]), \
             mock.patch('summarizer.paper_source.paper_text', return_value='Original English paper.'), \
             mock.patch.object(generated, 'chat_json', return_value=answer):
            with self.assertRaises(generated.GenerationError):
                slides.write('2101.00001v1', 'fr')(0)
