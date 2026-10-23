"""Native, editable PowerPoint objects from validated Journal Club pages."""
import hashlib
import base64
import io
import math
import re
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urljoin, urlsplit

import requests
from django.core.cache import cache
from django.utils.translation import gettext as _, get_language
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from . import sources

MAX_IMAGE_BYTES = 4 * 1024 * 1024
MAX_IMAGE_PIXELS = 12_000_000
MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


class ExportError(Exception):
    pass


def safe_image_url(url):
    try:
        parsed = urlsplit(url)
        return (parsed.scheme == 'https' and parsed.hostname in ('arxiv.org', 'www.arxiv.org', 'export.arxiv.org')
                and parsed.port in (None, 443) and not parsed.username and not parsed.password and parsed.path.startswith('/html/'))
    except ValueError:
        return False


def svg_png(data, width_limit=1800, height_limit=1200):
    """No scripts, external references, DTDs or local files reach the SVG renderer."""
    import resvg_py
    root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
    if root.getroottree().docinfo.doctype or etree.QName(root).localname != 'svg':
        raise ExportError
    for node in list(root.iter()):
        if not isinstance(node.tag, str):
            continue
        if etree.QName(node).localname in ('image', 'feImage'):
            href = next((v for k, v in node.attrib.items() if etree.QName(k).localname == 'href'), '')
            if not re.match(r'^data:image/(png|jpeg);base64,', href, flags=re.I):
                raise ExportError
            raster = base64.b64decode(href.split(',', 1)[1], validate=True)
            if len(raster) > MAX_IMAGE_BYTES or not raster.startswith((b'\x89PNG\r\n\x1a\n', b'\xff\xd8')):
                raise ExportError
            for attr in node.attrib:
                if etree.QName(attr).localname == 'href':
                    node.set(attr, 'data:image/png;base64,' + base64.b64encode(normalized_image(raster)).decode('ascii'))
        if etree.QName(node).localname in ('script', 'foreignObject'):
            if node.getparent() is not None:
                node.getparent().remove(node)
            continue
        for attr, value in list(node.attrib.items()):
            external_url = any(not target.strip(' \"\'').startswith('#') for target in re.findall(r'url\(([^)]*)\)', value, flags=re.I))
            inline_image = etree.QName(node).localname in ('image', 'feImage') and value.startswith('data:image/png;base64,')
            if (etree.QName(attr).localname in ('href', 'src') and not value.startswith('#') and not inline_image) or external_url or etree.QName(attr).localname.lower().startswith('on'):
                del node.attrib[attr]
        if etree.QName(node).localname == 'style':
            node.text = re.sub(r'@import[^;]*;', '', node.text or '', flags=re.I)
            node.text = re.sub(r'url\(([^)]*)\)', lambda match: match[0] if match[1].strip(' \"\'').startswith('#') else 'none', node.text, flags=re.I)
    box = (root.get('viewBox') or '').replace(',', ' ').split()
    try:
        width, height = map(float, box[2:]) if len(box) == 4 else (
            float(re.match(r'[\d.]+', root.get('width', '800')).group()),
            float(re.match(r'[\d.]+', root.get('height', '600')).group()))
        ratio = width / height
        if not math.isfinite(ratio) or not .05 <= ratio <= 20:
            raise ValueError
    except (ValueError, ZeroDivisionError, AttributeError):
        raise ExportError
    width, height = ((width_limit, round(width_limit / ratio)) if ratio >= width_limit / height_limit
                     else (round(height_limit * ratio), height_limit))
    return resvg_py.svg_to_bytes(svg_string=etree.tostring(root, encoding='unicode'), width=width, height=height)


def normalized_image(data):
    if data.lstrip().startswith(b'<'):
        data = svg_png(data)
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > MAX_IMAGE_PIXELS:
            raise ExportError
        image.load()
        image.thumbnail((1800, 1200))
        # A plain raster image part, without EXIF, active content or external relationships.
        image = image.convert('RGBA' if image.mode in ('RGBA', 'LA', 'P') else 'RGB')
        output = io.BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()


def figure_bytes(url):
    if not safe_image_url(url):
        raise ExportError
    key = 'pptx-figure:v1:' + hashlib.sha256(url.encode()).hexdigest()
    cached = cache.get(key)
    if cached is not None:
        return cached
    deadline = time.monotonic() + 7
    try:
        for _attempt in range(3):
            if not safe_image_url(url) or time.monotonic() > deadline:
                raise ExportError
            with requests.get(url, headers={'User-Agent': sources.USER_AGENT}, timeout=(2, 4),
                              stream=True, allow_redirects=False) as response:
                if 300 <= response.status_code < 400:
                    url = urljoin(url, response.headers.get('Location', ''))
                    continue
                response.raise_for_status()
                data = bytearray()
                for chunk in response.iter_content(65536):
                    data.extend(chunk)
                    if len(data) > MAX_IMAGE_BYTES or time.monotonic() > deadline:
                        raise ExportError
                result = normalized_image(bytes(data))
                cache.set(key, result, 6 * 3600)
                return result
    except (requests.RequestException, ValueError, OSError, RuntimeError, etree.XMLSyntaxError) as error:
        raise ExportError from error
    raise ExportError


def clean_text(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)


def text_size(lines, rect, maximum, minimum=14, bullets=False):
    """Conservative metrics, with slack for Office's substituted fonts. Never truncate text."""
    width, height = rect[2:]
    for size in range(round(maximum), minimum - 1, -1):
        font = ImageFont.load_default(size=round(size * 96 / 72))
        limit = max(1, width * 96 * .88 - (size * 1.5 if bullets else 0))
        count = 0
        for text in lines:
            length = 0
            for word in (text or ' ').split(' '):
                measured = font.getlength(word + ' ')
                if length and length + measured > limit:
                    count += 1
                    length = 0
                count += max(0, math.ceil(measured / limit) - 1)
                length += measured % limit if measured > limit else measured
            count += 1
        if count * size * 1.25 + max(0, len(lines) - 1) * 8 <= height * 72 * .92:
            return size
    raise ExportError(_('Some text does not fit. Shorten it or choose a smaller text size before exporting.'))


def textbox(slide, name, lines, rect, size, *, bold=False, centered=False, bullets=False, muted=False):
    x, y, width, height = rect
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.name = name
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if centered else MSO_ANCHOR.TOP
    for i, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = clean_text(text)
        paragraph.font._rPr.set('lang', 'fr-FR' if get_language() == 'fr' else 'en-US')
        paragraph.font.name = '+mj-lt' if bold else '+mn-lt'
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.theme_color = MSO_THEME_COLOR.DARK_2 if muted else MSO_THEME_COLOR.DARK_1
        paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(8 if i < len(lines) - 1 else 0)
        if bullets:
            props = paragraph._p.get_or_add_pPr()
            props.set('marL', str(Pt(size * 1.1)))
            props.set('indent', str(-Pt(size * .65)))
            marker = OxmlElement('a:buChar')
            marker.set('char', '•')
            props.append(marker)
    return box


def body_rects(page, rect):
    x, y, width, height = rect
    share, gap = page['figure_size'] / 100, .3
    if page['layout'] in ('left', 'right'):
        figure_width = width * share
        text_width = width - figure_width - gap
        if page['layout'] == 'left':
            return (x + figure_width + gap, y, text_width, height), (x, y, figure_width, height)
        return (x, y, text_width, height), (x + text_width + gap, y, figure_width, height)
    figure_height = height * share
    text_height = height - figure_height - gap
    if page['layout'] == 'top':
        return (x, y + figure_height + gap, width, text_height), (x, y, width, figure_height)
    return (x, y, width, text_height), (x, y + text_height + gap, width, figure_height)


def render_pptx(pages, arxiv_id, license_label, title):
    urls = list(dict.fromkeys(p['figure']['url'] for p in pages if p.get('figure')))
    try:
        with ThreadPoolExecutor(max_workers=4) as pool:
            images = dict(zip(urls, pool.map(figure_bytes, urls)))
    except ExportError:
        raise ExportError(_('A figure could not be included. Please try the PowerPoint export again later.'))
    presentation = Presentation()
    theme = presentation.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for family in ('majorFont', 'minorFont'):
        root.find('.//' + ns + family + '/' + ns + 'latin').set('typeface', 'Arial')
    for key, color in (('dk1', '172033'), ('dk2', '647084'), ('accent1', '5546D9')):
        entry = root.find('.//' + ns + 'clrScheme/' + ns + key)
        entry.clear()
        etree.SubElement(entry, ns + 'srgbClr', val=color)
    theme._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    presentation.slide_width = Inches(13 + 1 / 3)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = clean_text(title)
    presentation.core_properties.author = 'SummarizePaper'
    presentation.core_properties.language = get_language() or 'en'
    for number, page in enumerate(pages, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        is_title = page['kind'] == 'title'
        title_rect = (.58, 2 if is_title else .5, 12.17, 1.65 if is_title else 1.05)
        size = text_size([page['title']], title_rect, 36 if is_title else 30, minimum=22)
        textbox(slide, 'Slide title', [page['title']], title_rect, size, bold=True, centered=is_title)
        rect = (.58, 3.9, 12.17, 1.25) if is_title else (.58, 1.85, 12.17, 4.9)
        maximum = page['text_size'] * .75
        if page['kind'] == 'columns':
            width = (rect[2] - .5) / 2
            for i, (heading, lines) in enumerate(((_('Strengths'), page['bullets']), (_('Limits'), page['secondary']))):
                column = (rect[0] + i * (width + .5), rect[1] + .6, width, rect[3] - .6)
                textbox(slide, heading, [heading], (column[0], rect[1], width, .45), 22, bold=True)
                textbox(slide, heading + ' text', lines, column, text_size(lines, column, maximum, bullets=True), bullets=True)
        else:
            text_rect, figure_rect = body_rects(page, rect) if page.get('figure') else (rect, None)
            textbox(slide, 'Slide text', page['bullets'], text_rect,
                    text_size(page['bullets'], text_rect, maximum, bullets=not is_title), centered=is_title, bullets=not is_title)
            if figure_rect:
                data = images[page['figure']['url']]
                x, y, width, height = figure_rect
                caption_height = .52
                with Image.open(io.BytesIO(data)) as image:
                    factor = min(width / image.width, (height - caption_height) / image.height)
                    picture_width, picture_height = image.width * factor, image.height * factor
                picture = slide.shapes.add_picture(io.BytesIO(data), Inches(x + (width - picture_width) / 2),
                    Inches(y + (height - caption_height - picture_height) / 2), Inches(picture_width), Inches(picture_height))
                picture.name = 'Figure %s' % page['figure']['number']
                caption = _('Figure %(number)s of the paper (%(license)s)') % {'number': page['figure']['number'], 'license': license_label}
                textbox(slide, 'Figure credit', [caption], (x, y + height - caption_height, width, caption_height), 10, centered=True, muted=True)
        textbox(slide, 'Source', ['SummarizePaper · ' + sources.label(arxiv_id)], (.58, 7.04, 10.8, .24), 10, muted=True)
        textbox(slide, 'Slide number', ['%s / %s' % (number, len(pages))], (11.8, 7.04, .95, .24), 10, muted=True)
        if page.get('notes'):
            notes = slide.notes_slide.notes_text_frame
            notes.text = clean_text(page['notes'])
            for paragraph in notes.paragraphs:
                paragraph.font._rPr.set('lang', 'fr-FR' if get_language() == 'fr' else 'en-US')
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
