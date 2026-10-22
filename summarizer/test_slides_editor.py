import copy
import json
import io
from unittest import mock

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from django.test import Client

from . import slides
from . import slides_pptx
from .models import Generated, UserPlan
from .test_tools import ToolsTestCase


class SlidesEditorTests(ToolsTestCase):
    def setUp(self):
        super().setUp()
        self.article = self.paper()
        self.reader = self.user('editor', paid=True)
        self.client = self.client_for(self.reader)
        self.url = '/en/journal-club/2101.00001v1/'
        self.data = {'slides': [{'title': 'Results', 'bullets': ['Original result'], 'notes': 'Private speaker note',
                                'figure': {'number': 1, 'url': 'https://arxiv.org/html/2101.00001v1/figure.png', 'caption': 'Spectrum'}}],
                     'strengths': ['Evidence'], 'limitations': ['Small sample'], 'questions': ['Why?'], 'license': 'CC BY 4.0'}
        self.original = Generated.objects.create(kind='slides', key=self.article.arxiv_id, lang='en', status='done', content=json.dumps(self.data))

    def payload(self, client=None):
        response = (client or self.client).get(self.url)
        return {'revision': response.context['deck_revision'], 'pages': copy.deepcopy(response.context['pages'])}

    def save(self, payload, client=None):
        return (client or self.client).post(self.url + 'save/', json.dumps(payload), content_type='application/json')

    def export(self, payload, client=None):
        return (client or self.client).post(self.url + 'powerpoint/', json.dumps(payload), content_type='application/json')

    def image(self):
        output = io.BytesIO()
        Image.new('RGB', (640, 400), 'white').save(output, 'PNG')
        return output.getvalue()

    def test_powerpoint_contains_editable_text_picture_notes_and_unsaved_changes(self):
        payload = self.payload()
        payload['pages'][1].update(title='My unsaved interpretation', layout='left', figure_size=60, notes='Speaker-only note')
        payload['pages'][1]['figure'] = {'url': 'http://127.0.0.1/private'}
        with mock.patch.object(slides_pptx, 'figure_bytes', return_value=self.image()) as load:
            response = self.export(payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], slides_pptx.MIME)
        self.assertIn('.pptx', response['Content-Disposition'])
        self.assertIn('no-store', response['Cache-Control'])
        load.assert_called_once_with(self.data['slides'][0]['figure']['url'])
        deck = Presentation(io.BytesIO(response.content))
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9)
        self.assertEqual(len(deck.slides), 4)
        for page in deck.slides:
            for shape in page.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
        slide = deck.slides[1]
        boxes = {shape.name: shape for shape in slide.shapes}
        self.assertEqual(boxes['Slide title'].text, 'My unsaved interpretation')
        self.assertEqual(boxes['Figure 1'].shape_type, MSO_SHAPE_TYPE.PICTURE)
        self.assertLess(boxes['Figure 1'].left, boxes['Slide text'].left)
        self.assertIn('CC BY 4.0', boxes['Figure credit'].text)
        self.assertEqual(slide.notes_slide.notes_text_frame.text, 'Speaker-only note')
        self.assertNotIn('Speaker-only note', '\n'.join(s.text for s in slide.shapes if s.has_text_frame))
        self.assertFalse(Generated.objects.filter(kind='slides_edit').exists())
        self.original.refresh_from_db()
        self.assertEqual(self.original.data(), self.data)

    def test_powerpoint_is_private_and_requires_premium_csrf_and_post(self):
        payload = self.payload()
        with mock.patch.object(slides_pptx, 'render_pptx') as render:
            self.assertEqual(self.client.get(self.url + 'powerpoint/').status_code, 405)
            self.assertEqual(self.export(payload, Client()).status_code, 403)
            strict = Client(enforce_csrf_checks=True)
            strict.force_login(self.reader)
            self.assertEqual(self.export(payload, strict).status_code, 403)
            UserPlan.objects.filter(user=self.reader).update(paid=False)
            self.assertEqual(self.export(payload).status_code, 403)
            render.assert_not_called()

    def test_powerpoint_checks_revision_input_and_missing_generated_deck(self):
        payload = self.payload()
        with mock.patch.object(slides_pptx, 'render_pptx') as render:
            self.assertEqual(self.export(dict(payload, revision='stale')).status_code, 409)
            payload['pages'][1]['layout'] = 'invalid'
            self.assertEqual(self.export(payload).status_code, 400)
            self.original.status = 'pending'
            self.original.save(update_fields=['status'])
            self.assertEqual(self.export(payload).status_code, 404)
            render.assert_not_called()

    def test_failed_image_fetch_returns_a_retryable_error_instead_of_an_incomplete_deck(self):
        payload = self.payload()
        with mock.patch.object(slides_pptx, 'figure_bytes', side_effect=slides_pptx.ExportError):
            response = self.export(payload)
        self.assertEqual(response.status_code, 503)
        self.assertIn('figure could not be included', response.json()['message'])
        self.assertFalse(Generated.objects.filter(kind='slides_edit').exists())

    def test_untrusted_image_urls_and_redirects_are_never_fetched(self):
        with mock.patch.object(slides_pptx.requests, 'get') as get:
            for url in ('http://arxiv.org/html/id/a.png', 'https://arxiv.org.evil.test/html/a.png',
                        'https://user@arxiv.org/html/a.png', 'https://arxiv.org:bad/html/a.png', 'http://127.0.0.1/a.png'):
                with self.assertRaises(slides_pptx.ExportError):
                    slides_pptx.figure_bytes(url)
            get.assert_not_called()
            response = get.return_value.__enter__.return_value
            response.status_code = 302
            response.headers = {'Location': 'http://127.0.0.1/private'}
            with self.assertRaises(slides_pptx.ExportError):
                slides_pptx.figure_bytes('https://arxiv.org/html/id/a.png')
            self.assertEqual(get.call_count, 1)

    def test_svg_figures_keep_local_gradients_without_external_resources(self):
        svg = b'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 300"><defs><linearGradient id="g"><stop stop-color="blue"/></linearGradient></defs><rect width="400" height="300" fill="url(#g)"/></svg>'''
        raster = slides_pptx.svg_png(svg)
        self.assertTrue(raster.startswith(b'\x89PNG'))
        with self.assertRaises(slides_pptx.ExportError):
            slides_pptx.svg_png(b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg xmlns="http://www.w3.org/2000/svg"/>')
        with self.assertRaises(slides_pptx.ExportError):
            slides_pptx.svg_png(b'<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///etc/passwd"/></svg>')

    def test_figures_are_cached_and_size_limits_reject_oversized_downloads(self):
        with mock.patch.object(slides_pptx.requests, 'get') as get:
            response = get.return_value.__enter__.return_value
            response.status_code = 200
            response.iter_content.return_value = [self.image()]
            url = 'https://arxiv.org/html/id/cache.png'
            first = slides_pptx.figure_bytes(url)
            self.assertTrue(first.startswith(b'\x89PNG'))
            self.assertEqual(slides_pptx.figure_bytes(url), first)
            self.assertEqual(get.call_count, 1)
            response.iter_content.return_value = [b'x' * (slides_pptx.MAX_IMAGE_BYTES + 1)]
            with self.assertRaises(slides_pptx.ExportError):
                slides_pptx.figure_bytes('https://arxiv.org/html/id/too-large.png')

    def test_powerpoint_rejects_text_that_cannot_fit_without_truncating_it(self):
        payload = self.payload()
        payload['pages'][1]['bullets'] = ['Long scientific explanation ' * 20] * 8
        with mock.patch.object(slides_pptx, 'figure_bytes', return_value=self.image()):
            response = self.export(payload)
        self.assertEqual(response.status_code, 503)
        self.assertIn('does not fit', response.json()['message'])

    def test_text_layout_and_notes_are_saved_privately_without_changing_original(self):
        payload = self.payload()
        payload['pages'][1].update(title='My interpretation', bullets=['A new result'], notes='My notes', layout='left', figure_size=65)
        result = self.save(payload)
        self.assertEqual(result.status_code, 200)
        page = self.client.get(self.url)
        self.assertContains(page, 'My interpretation')
        self.assertContains(page, 'data-layout="left"')
        self.assertContains(page, '--figure-share:65%')
        self.assertContains(page, 'My notes')
        self.assertEqual(page.context['deck_revision'], result.json()['revision'])
        self.assertIn('no-store', page['Cache-Control'])
        other = self.client_for(self.user('other-editor', paid=True))
        self.assertNotContains(other.get(self.url), 'My interpretation')
        self.original.refresh_from_db()
        self.assertEqual(self.original.data(), self.data)
        private = Generated.objects.get(kind='slides_edit')
        self.assertEqual(private.user_id, self.reader.pk)
        self.assertEqual(other.get('/en/ai/%s/' % private.pk).status_code, 404)

    def test_stale_tab_cannot_silently_overwrite_saved_changes(self):
        old = self.payload()
        changed = copy.deepcopy(old)
        changed['pages'][0]['title'] = 'A changed title'
        self.assertEqual(self.save(changed).status_code, 200)
        self.assertEqual(self.save(old).status_code, 409)
        self.assertContains(self.client.get(self.url), 'A changed title')

    def test_client_cannot_replace_source_images_or_credits_and_html_is_text(self):
        payload = self.payload()
        payload['pages'][1]['figure'] = {'url': 'http://127.0.0.1/private', 'number': 999}
        payload['pages'][1]['title'] = '<script>alert(1)</script>'
        self.assertEqual(self.save(payload).status_code, 200)
        page = self.client.get(self.url)
        self.assertContains(page, self.data['slides'][0]['figure']['url'])
        self.assertNotContains(page, 'http://127.0.0.1/private')
        self.assertContains(page, 'Figure 1 of the paper (arXiv, CC BY 4.0)')
        self.assertNotContains(page, '<script>alert(1)</script>')
        self.assertContains(page, '&lt;script&gt;alert(1)&lt;/script&gt;')

    def test_invalid_editor_input_is_rejected_without_creating_a_copy(self):
        for field, value in [('layout', 'absolute'), ('figure_size', '50;position:fixed'), ('text_size', 200),
                             ('title', ''), ('bullets', ['x'] * 9), ('notes', 'x' * 3001)]:
            payload = self.payload()
            payload['pages'][1][field] = value
            self.assertEqual(self.save(payload).status_code, 400)
        payload = self.payload()
        payload['pages'].pop()
        self.assertEqual(self.save(payload).status_code, 400)
        self.assertFalse(Generated.objects.filter(kind='slides_edit').exists())

    def test_save_requires_paid_account_csrf_and_post(self):
        payload = self.payload()
        self.assertEqual(self.client.get(self.url + 'save/').status_code, 405)
        self.assertEqual(self.save(payload, Client()).status_code, 403)
        strict = Client(enforce_csrf_checks=True)
        strict.force_login(self.reader)
        self.assertEqual(self.save(payload, strict).status_code, 403)
        UserPlan.objects.filter(user=self.reader).update(paid=False)
        self.assertEqual(self.save(payload).status_code, 403)

    def test_print_and_editor_controls_preserve_generation_and_source_warning(self):
        page = self.client.get(self.url)
        self.assertContains(page, 'slides.css')
        self.assertContains(page, 'slides.js')
        self.assertContains(page, 'Save as PDF</button>', count=0)
        self.assertContains(page, '16:9')
        self.assertContains(page, 'Speaker notes (excluded from PDF)')
        self.assertContains(page, 'data-edit="layout"')
        self.assertContains(page, 'data-edit="figure_size"')
        self.original.status = 'pending'
        self.original.save(update_fields=['status'])
        self.assertNotContains(self.client.get(self.url), 'id="deck-edit"')
